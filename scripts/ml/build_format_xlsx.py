"""학습용 표의 **형식**을 엑셀로 낸다 — 1번 시트 예시, 2번 시트 칸 설명.

왜 엑셀인가
  칸을 더할지 뺄지 사람이 보고 정해야 한다. CSV 는 칸 이름만 보이고 뜻이 안 보인다.
  1번 시트에서 실제 줄 모양을 보고, 2번 시트에서 그 칸이 무엇인지 읽는다.

⚠️ 여기 들어가는 줄은 **가상 예시**다. 학습에 쓰지 않는다.
   실제 데이터는 나중에 data/ml/ 아래에 만들어지고, 그 폴더는 git-ignore 다.
   이 폴더(docs/ml_format)는 **형식과 판단**을 담는 곳이라 커밋한다 —
   처음에 형식 파일을 data/ml 에 두었더니 git 이 통째로 무시해 "바뀐 게 없다"로 보였다.

⚠️ 값이 비어 있거나 status 가 '후보/보류/제외권장' 인 것은 **아직 안 정한 것**이다.
   정하기 전에 데이터를 채우면 나중에 표본을 다시 만들어야 한다.

산출: docs/ml_format/*.xlsx

⚠️ CSV 는 BOM(utf-8-sig)으로 쓴다. 엑셀이 Windows 에서 BOM 없는 UTF-8 을 cp949 로 읽어
   한글이 깨진다. 읽을 때는 utf-8-sig 로 벗긴다.
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
#  폴더를 인자로 받는다 — 학습표(docs/ml_format)와 핵심 스키마(docs/schema)가 같은 도구를 쓴다.
D = REPO / (sys.argv[1] if len(sys.argv) > 1 else "docs/ml_format")
E = D / "examples"

HEAD_FILL = PatternFill("solid", fgColor="E8E5E0")
NOTE_FILL = PatternFill("solid", fgColor="FBF3E2")     # 아직 안 정한 칸
DROP_FILL = PatternFill("solid", fgColor="F6E2E2")     # 넣지 말자고 적어 둔 칸
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def fit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_head(ws, row):
    ws.append(row)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = TOP
    ws.freeze_panes = "A2"


def run():
    cat = list(csv.DictReader((D / "_columns.csv").open(encoding="utf-8-sig")))
    #  표 목록도 카탈로그에서 읽는다. 코드에 적어 두면 표를 더할 때마다 고쳐야 한다.
    tables = list(dict.fromkeys(c["file"] for c in cat))

    for t in tables:
        cols = [c for c in cat if c["file"] == t]
        head = [c["column"] for c in cols if c["status"] != "제외권장"]
        rows = list(csv.DictReader((E / f"{t}.csv").open(encoding="utf-8-sig")))

        wb = Workbook()

        # ── 1번 시트: 실제 줄 모양 (가상 예시) ──────────────────────────
        ws = wb.active
        ws.title = "데이터"
        sheet_head(ws, head)
        for r in rows:
            ws.append([r.get(k, "") for k in head])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = TOP
        fit(ws, [max(12, min(34, len(h) + 8)) for h in head])

        # ── 2번 시트: 이 칸이 무엇인가 ────────────────────────────────
        ws2 = wb.create_sheet("칸 설명")
        sheet_head(ws2, ["칸", "형", "필수", "허용값", "어디서 오나",
                         "상태", "뜻", "예시"])
        for c in cols:
            #  예시는 1번 시트에서 **실제로 채워진 값**을 가져온다.
            #  설명과 예시가 따로 놀면 칸을 판단할 수 없다.
            ex = next((r[c["column"]] for r in rows
                       if r.get(c["column"], "").strip()), "")
            ws2.append([c["column"], c["type"], c["required"], c["allowed"],
                        c["source"], c["status"], c["note"], ex])
            last = ws2[ws2.max_row]
            for cell in last:
                cell.alignment = WRAP
            if c["status"] == "제외권장":
                for cell in last:
                    cell.fill = DROP_FILL
            elif c["status"] in ("후보", "보류"):
                for cell in last:
                    cell.fill = NOTE_FILL
        fit(ws2, [20, 8, 6, 26, 14, 10, 62, 34])
        ws2.auto_filter.ref = ws2.dimensions

        #  헤더 CSV 도 여기서 만든다. 손으로 두 곳을 맞추면 반드시 어긋난다.
        #  BOM 을 붙여야 엑셀이 cp949 로 읽지 않는다.
        with (D / f"{t}.csv").open("w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh).writerow(head)

        out = D / f"{t}.xlsx"
        wb.save(out)
        n_open = sum(1 for c in cols if c["status"] in ("후보", "보류", "제외권장"))
        print(f"  {t + '.xlsx':22} 칸 {len(head)}개 · 예시 {len(rows)}줄"
              + (f" · 판단 대기 {n_open}칸" if n_open else ""))

    # ── 설정도 같은 모양으로 (config.csv 가 있는 폴더에서만) ──────────────
    if not (D / "config.csv").exists():
        print()
        print(f"  → {D.relative_to(REPO)}"
              "  (노란 줄 = 아직 안 정한 것, 붉은 줄 = 빼자고 적어 둔 것)")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "설정"
    sheet_head(ws, ["key", "value", "type", "allowed", "note"])
    todo = 0
    for r in csv.DictReader((D / "config.csv").open(encoding="utf-8-sig")):
        ws.append([r["key"], r["value"], r["type"], r["allowed"], r["note"]])
        last = ws[ws.max_row]
        for cell in last:
            cell.alignment = WRAP
        if not r["value"]:
            todo += 1
            for cell in last:
                cell.fill = NOTE_FILL          # 값이 비었다 = 아직 안 정했다
    fit(ws, [22, 18, 10, 24, 70])

    ws2 = wb.create_sheet("칸 설명")
    sheet_head(ws2, ["칸", "뜻", "예시"])
    for k, v, ex in [
        ("key", "무엇을 정하는 설정인가", "companyCap.mode"),
        ("value", "정한 값. **비어 있으면 아직 안 정한 것**이다", "weight"),
        ("type", "값의 형", "enum"),
        ("allowed", "고를 수 있는 값. | 로 잇는다", "none|drop|weight"),
        ("note", "왜 정해야 하는가 · 지금 자료가 어떤 상태인가",
         "쿠팡 296건(코퍼스의 12%) 처리"),
    ]:
        ws2.append([k, v, ex])
        for cell in ws2[ws2.max_row]:
            cell.alignment = WRAP
    fit(ws2, [14, 70, 26])
    wb.save(D / "config.xlsx")
    print(f"  {'config.xlsx':22} 항목 {ws.max_row - 1}개 · 값 미정 {todo}개")

    print(f"\n  → {D.relative_to(REPO)}  (노란 줄 = 아직 안 정한 것, 붉은 줄 = 빼자고 적어 둔 것)")


if __name__ == "__main__":
    run()
